import argparse
import uuid
from pathlib import Path

import psycopg2
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# ---------------- CONFIG ----------------

WORDS_PER_SEGMENT = 10

DB_CONFIG = {
    "dbname": "vaanisetu",
    "user": "postgres",
    "password": "akinikar06",  # change this
    "host": "localhost",
    "port": "5432"
}

HEADING_STYLES = {
    "heading 1", "heading 2", "heading 3",
    "heading 4", "heading 5", "heading 6"
}
LIST_STYLES = {
    "list paragraph", "listbullet", "listnumber",
    "list bullet", "list number"
}
CAPTION_STYLES = {"caption"}

# ---------------- DB FUNCTIONS ----------------

def get_connection():
    return psycopg2.connect(**DB_CONFIG)


def insert_document(project_id, filename):
    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
        INSERT INTO documents (project_id, filename, s3_key, status)
        VALUES (%s, %s, %s, %s)
        RETURNING id;
    """, (project_id, filename, "local_upload", "PARSED"))

    doc_id = cur.fetchone()[0]

    conn.commit()
    cur.close()
    conn.close()

    return doc_id


def insert_segment(document_id, target_lang, text):
    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
        INSERT INTO segments (
            document_id,
            target_lang,
            source_text,
            translated_text,
            translation_source
        )
        VALUES (%s, %s, %s, %s, %s);
    """, (
        document_id,
        target_lang,
        text,
        None,
        None
    ))

    conn.commit()
    cur.close()
    conn.close()

# ---------------- HELPERS ----------------

def _style_name(para):
    try:
        return (para.style.name or "").lower().strip()
    except:
        return ""


def _is_list_para(para):
    if _style_name(para) in LIST_STYLES:
        return True
    pPr = para._p.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:numPr")) is not None


def split_text(text, words_per_segment):
    words = text.split()
    return [
        " ".join(words[i:i + words_per_segment])
        for i in range(0, len(words), words_per_segment)
    ]

# ---------------- PARSER ----------------

def parse_and_store(input_path, project_id, target_lang):
    doc = Document(str(input_path))

    print("[parser] Parsing document...")

    # Step 1: insert document
    document_id = insert_document(project_id, input_path.name)
    print(f"[parser] Document ID: {document_id}")

    segment_count = 0

    # Step 2: iterate paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        chunks = split_text(text, WORDS_PER_SEGMENT)

        for chunk in chunks:
            insert_segment(document_id, target_lang, chunk)
            segment_count += 1

    # Step 3: iterate tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue

                    chunks = split_text(text, WORDS_PER_SEGMENT)

                    for chunk in chunks:
                        insert_segment(document_id, target_lang, chunk)
                        segment_count += 1

    print(f"[parser] Inserted {segment_count} segments into DB")

# ---------------- MAIN ----------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--project-id", type=int, default=1)
    parser.add_argument("--target-lang", default="en-US")

    args = parser.parse_args()

    input_path = Path(args.input)

    if not input_path.exists():
        raise FileNotFoundError("File not found")

    print(f"[parser] Reading: {input_path}")

    parse_and_store(
        input_path,
        args.project_id,
        args.target_lang
    )

    print("[parser] Done")


if __name__ == "__main__":
    main()