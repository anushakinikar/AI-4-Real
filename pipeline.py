"""
VaaniSetu — End-to-End Translation Pipeline
============================================
Steps run serially for a given .docx file:

    1. PARSING          — extract segments from .docx → documents + segments tables
    2. VALIDATION       — grammar/spelling check via LanguageTool → validation_issues table
    3. TRANSLATION      — EN → DE via Azure Cognitive Translator + TM lookup → segments table
    4. BACK-TRANSLATION — DE → EN via MyMemory + quality scoring → agent_segment_evaluation table

Usage (only --input is required):
    python pipeline.py --input "path/to/document.docx"

Requirements:
    pip install psycopg2-binary requests python-dotenv rapidfuzz sentence-transformers scikit-learn python-docx
"""

import argparse
import hashlib
import os
import sys
import time
from pathlib import Path
from typing import Optional

import requests
import psycopg2
import psycopg2.extras
from dotenv import load_dotenv

load_dotenv()

# ══════════════════════════════════════════════════════════════════════════════
# SHARED CONFIG
# ══════════════════════════════════════════════════════════════════════════════

DB_CONFIG = {
    "host":     os.getenv("DB_HOST",     "localhost"),
    "port":     int(os.getenv("DB_PORT", "5432")),
    "dbname":   os.getenv("DB_NAME",     "vaanisetu"),
    "user":     os.getenv("DB_USER",     "postgres"),
    "password": os.getenv("DB_PASSWORD", ""),#insert password here
}

AZURE_KEY      = os.getenv("AZURE_TRANSLATOR_KEY", "") #insert key here
AZURE_REGION   = os.getenv("AZURE_TRANSLATOR_REGION",   "centralindia")
AZURE_ENDPOINT = os.getenv("AZURE_TRANSLATOR_ENDPOINT", "https://api.cognitive.microsofttranslator.com")
AZURE_API_VER  = "3.0"
BATCH_SIZE     = 100
BATCH_CHAR_CAP = 50_000

MYMEMORY_URL   = "https://api.mymemory.translated.net/get"
MYMEMORY_EMAIL = None   # set to your email for higher rate limit

LT_API_URL    = "https://api.languagetool.org/v2/check"
LT_RATE_LIMIT = 3.1     # seconds between LanguageTool calls
LT_MAX_CHARS  = 19000

WORDS_PER_SEGMENT = 10


# ══════════════════════════════════════════════════════════════════════════════
# DB HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def get_conn():
    return psycopg2.connect(**DB_CONFIG)


def get_dict_conn():
    return psycopg2.connect(**DB_CONFIG, cursor_factory=psycopg2.extras.RealDictCursor)


def sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def banner(title: str):
    print("\n" + "=" * 65)
    print(f"  {title}")
    print("=" * 65)


def step_banner(n: int, title: str):
    print(f"\n{'─' * 65}")
    print(f"  STEP {n}: {title}")
    print(f"{'─' * 65}")


# ══════════════════════════════════════════════════════════════════════════════
# STEP 1 — PARSING
# Schema tables used: documents, segments
# ══════════════════════════════════════════════════════════════════════════════

def split_text(text: str, words_per_segment: int) -> list:
    words = text.split()
    return [
        " ".join(words[i:i + words_per_segment])
        for i in range(0, len(words), words_per_segment)
    ]


def run_parsing(input_path: Path, project_id: int, target_lang: str) -> int:
    """
    Parse .docx → inserts 1 row into documents, N rows into segments.
    Returns the new document_id (INT PK).

    documents columns used:
        project_id, filename, s3_key, status, target_lang, sensitivity

    segments columns used:
        document_id, target_lang, source_text, translated_text, translation_source
    """
    step_banner(1, "PARSING")
    from docx import Document

    print(f"  Reading : {input_path}")

    conn = get_conn()
    cur  = conn.cursor()

    # Insert document row — target_lang and sensitivity are real columns in schema
    cur.execute("""
        INSERT INTO documents (project_id, filename, s3_key, status, target_lang, sensitivity)
        VALUES (%s, %s, %s, %s::document_status, %s, %s::sensitivity_level)
        RETURNING id;
    """, (project_id, input_path.name, "local_upload", "PARSED", target_lang, "STANDARD"))
    document_id = cur.fetchone()[0]
    conn.commit()

    doc = Document(str(input_path))
    segment_count = 0

    def _insert_chunk(text_chunk: str):
        nonlocal segment_count
        cur.execute("""
            INSERT INTO segments
                (document_id, target_lang, source_text, translated_text, translation_source)
            VALUES (%s, %s, %s, NULL, NULL);
        """, (document_id, target_lang, text_chunk))
        segment_count += 1

    # Paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        for chunk in split_text(text, WORDS_PER_SEGMENT):
            _insert_chunk(chunk)

    # Table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    for chunk in split_text(text, WORDS_PER_SEGMENT):
                        _insert_chunk(chunk)

    conn.commit()
    cur.close()
    conn.close()

    print(f"  Document ID : {document_id}")
    print(f"  Segments    : {segment_count} inserted into DB")
    return document_id


# ══════════════════════════════════════════════════════════════════════════════
# STEP 2 — VALIDATION
# Schema table used: validation_issues
#   Columns: segment_id, document_id, type (validation_issue_type enum),
#            severity (severity_level enum), message, offset_start, offset_end
# ══════════════════════════════════════════════════════════════════════════════

# Maps LanguageTool category IDs → validation_issue_type enum values (from schema)
_ISSUE_TYPE_MAP = {
    "TYPOS":                  "SPELLING",
    "MISSPELLING":            "SPELLING",
    "GRAMMAR":                "GRAMMAR",
    "AGREEMENT":              "GRAMMAR",
    "SUBJECT_VERB_AGREEMENT": "GRAMMAR",
    "VERB_FORM":              "GRAMMAR",
    "PUNCTUATION":            "PUNCTUATION",
    "TYPOGRAPHY":             "PUNCTUATION",
    "CONSISTENCY":            "CONSISTENCY",
    "STYLE":                  "CONSISTENCY",
    "CASING":                 "FORMATTING",
    "FORMATTING":             "FORMATTING",
}

# Maps LanguageTool category IDs → severity_level enum values (from schema)
_SEVERITY_MAP = {
    "TYPOS":                  "ERROR",
    "MISSPELLING":            "ERROR",
    "GRAMMAR":                "ERROR",
    "AGREEMENT":              "ERROR",
    "SUBJECT_VERB_AGREEMENT": "ERROR",
    "VERB_FORM":              "ERROR",
    "PUNCTUATION":            "WARNING",
    "TYPOGRAPHY":             "WARNING",
    "CONSISTENCY":            "WARNING",
    "CASING":                 "WARNING",
    "STYLE":                  "INFO",
    "FORMATTING":             "WARNING",
}


def _map_issue_type(match: dict) -> str:
    category = match.get("rule", {}).get("category", {}).get("id", "").upper()
    return _ISSUE_TYPE_MAP.get(category, "GRAMMAR")


def _map_severity(match: dict) -> str:
    category = match.get("rule", {}).get("category", {}).get("id", "").upper()
    return _SEVERITY_MAP.get(category, "WARNING")


def _lt_check(text: str, lang: str) -> list:
    if not text.strip():
        return []
    try:
        r = requests.post(
            LT_API_URL,
            data={"text": text[:LT_MAX_CHARS], "language": lang},
            timeout=15,
        )
        r.raise_for_status()
        return r.json().get("matches", [])
    except Exception as e:
        print(f"  [warn] LanguageTool API error: {e}")
        return []


def run_validation(document_id: int, lang: str = "en-US") -> dict:
    """
    Validate source segments via LanguageTool API.
    Writes issues into schema's validation_issues table (already defined in schema).

    validation_issues columns written:
        segment_id, document_id, type, severity, message, offset_start, offset_end

    Returns: { segment_id: [issue_messages] }
    """
    step_banner(2, "VALIDATION")

    conn = get_conn()
    cur  = conn.cursor()

    cur.execute("""
        SELECT id, source_text
        FROM   segments
        WHERE  document_id = %s
        ORDER  BY id;
    """, (document_id,))
    rows  = cur.fetchall()
    total = len(rows)

    print(f"  Document ID : {document_id}")
    print(f"  Segments    : {total}")
    print(f"  Language    : {lang}\n")

    issues_map   = {}
    issues_total = 0

    for i, (seg_id, source_text) in enumerate(rows):
        print(f"  [{i+1}/{total}] seg {seg_id}: {source_text[:60]}")
        matches = _lt_check(source_text, lang)

        if not matches:
            print("           ✓ clean")
        else:
            msgs = []
            for m in matches:
                msg          = m.get("message", "")
                offset_start = m.get("offset", 0)
                offset_end   = offset_start + m.get("length", 0)
                wrong        = source_text[offset_start:offset_end]
                issue_type   = _map_issue_type(m)
                severity     = _map_severity(m)

                print(f"           ✗ [{severity}] {msg}  →  '{wrong}'")
                msgs.append(msg)

                # Write to the real validation_issues table from schema
                cur.execute("""
                    INSERT INTO validation_issues
                        (segment_id, document_id, type, severity, message, offset_start, offset_end)
                    VALUES (%s, %s, %s::validation_issue_type, %s::severity_level, %s, %s, %s);
                """, (seg_id, document_id, issue_type, severity, msg, offset_start, offset_end))
                issues_total += 1

            issues_map[seg_id] = msgs

        if i < total - 1:
            time.sleep(LT_RATE_LIMIT)

    conn.commit()
    cur.close()
    conn.close()

    print(f"\n  Total segments      : {total}")
    print(f"  Segments with issues: {len(issues_map)}")
    print(f"  Total issues found  : {issues_total}")
    print(f"  Issues logged to    : validation_issues table")

    return issues_map


# ══════════════════════════════════════════════════════════════════════════════
# STEP 3 — ACTUAL TRANSLATION (Azure Cognitive Translator + TM)
# Schema tables used: segments (UPDATE), tm_entries (SELECT + INSERT/UPDATE)
#
# NOTE: tm_entries.source_hash has an INDEX but NO UNIQUE constraint in schema,
#       so ON CONFLICT (source_hash) fails. We use SELECT + INSERT/UPDATE instead.
# ══════════════════════════════════════════════════════════════════════════════

def _azure_batch(texts: list, target_lang: str, source_lang: str) -> list:
    if not texts:
        return []
    url = f"{AZURE_ENDPOINT.rstrip('/')}/translate"
    headers = {
        "Ocp-Apim-Subscription-Key":    AZURE_KEY,
        "Ocp-Apim-Subscription-Region": AZURE_REGION,
        "Content-Type":                 "application/json",
    }
    params = {"api-version": AZURE_API_VER, "from": source_lang, "to": target_lang}
    body   = [{"text": t} for t in texts]

    resp = requests.post(url, headers=headers, params=params, json=body, timeout=30)
    resp.raise_for_status()

    return [
        item.get("translations", [{}])[0].get("text", "")
        for item in resp.json()
    ]


def _azure_translate_in_batches(texts: list, target_lang: str, source_lang: str) -> list:
    all_results = []
    batch       = []
    batch_chars = 0

    def flush(b):
        if b:
            all_results.extend(_azure_batch(b, target_lang, source_lang))
            time.sleep(0.2)

    for text in texts:
        if batch and (len(batch) >= BATCH_SIZE or batch_chars + len(text) > BATCH_CHAR_CAP):
            flush(batch)
            batch       = []
            batch_chars = 0
        batch.append(text)
        batch_chars += len(text)

    flush(batch)
    return all_results


def _upsert_tm_entry(cur, org_id, source_lang, target_lang, src_text, tgt_text, h, domain):
    """
    Safe TM upsert without ON CONFLICT — schema has no UNIQUE constraint on source_hash.
    Checks for existing row first, then updates usage_count or inserts fresh row.
    """
    cur.execute("""
        SELECT id FROM tm_entries
        WHERE source_hash = %s AND source_lang = %s AND target_lang = %s AND org_id = %s
        LIMIT 1;
    """, (h, source_lang, target_lang, org_id))
    existing = cur.fetchone()

    if existing:
        existing_id = existing["id"] if isinstance(existing, dict) else existing[0]
        cur.execute("""
            UPDATE tm_entries
            SET usage_count = usage_count + 1,
                target_text = %s
            WHERE id = %s;
        """, (tgt_text, existing_id))
    else:
        cur.execute("""
            INSERT INTO tm_entries
                (org_id, source_lang, target_lang, source_text,
                 target_text, source_hash, domain, usage_count)
            VALUES (%s, %s, %s, %s, %s, %s, %s, 1);
        """, (org_id, source_lang, target_lang, src_text, tgt_text, h, domain))


def run_translation(
    document_id: int,
    target_lang: str,
    source_lang: str       = "en",
    org_id: int            = 1,
    domain: Optional[str]  = None,
    dry_run: bool          = False,
):
    """
    Translate all segments for document_id using TM exact-match + Azure.

    segments columns updated:
        translated_text, translation_source (TM_EXACT or LLM — from translation_source enum)

    tm_entries columns written:
        org_id, source_lang, target_lang, source_text, target_text,
        source_hash, domain, usage_count
    """
    step_banner(3, "ACTUAL TRANSLATION")

    conn = get_dict_conn()
    cur  = conn.cursor()

    cur.execute("""
        SELECT id, source_text
        FROM   segments
        WHERE  document_id = %s
        ORDER  BY id;
    """, (document_id,))
    rows  = list(cur.fetchall())
    total = len(rows)

    print(f"  Document ID  : {document_id}")
    print(f"  Segments     : {total}")
    print(f"  {source_lang}  →  {target_lang}")
    print(f"  Dry run      : {dry_run}\n")

    if total == 0:
        print("  [warn] No segments found. Skipping translation.")
        cur.close()
        conn.close()
        return

    # TM exact-match lookup by SHA-256 hash
    seg_texts = [r["source_text"] for r in rows]
    hashes    = {sha256(t): t for t in seg_texts}

    cur.execute("""
        SELECT source_hash, target_text
        FROM   tm_entries
        WHERE  source_hash = ANY(%s)
          AND  source_lang = %s
          AND  target_lang = %s
          AND  org_id      = %s;
    """, (list(hashes.keys()), source_lang, target_lang, org_id))

    tm_hits = {hashes[r["source_hash"]]: r["target_text"] for r in cur.fetchall()}
    print(f"  TM hits      : {len(tm_hits)} / {total}")

    # Cache misses go to Azure
    misses = [(r["id"], r["source_text"]) for r in rows if r["source_text"] not in tm_hits]
    print(f"  Azure needed : {len(misses)}\n")

    azure_map = {}   # seg_id → translated_text

    if misses and not dry_run:
        miss_ids   = [m[0] for m in misses]
        miss_texts = [m[1] for m in misses]
        print(f"  Calling Azure Translator for {len(misses)} segments …")
        try:
            translated = _azure_translate_in_batches(miss_texts, target_lang, source_lang)
        except Exception as exc:
            print(f"  [ERROR] Azure translation failed: {exc}")
            cur.close()
            conn.close()
            raise

        for seg_id, src, tgt in zip(miss_ids, miss_texts, translated):
            azure_map[seg_id] = tgt
            print(f"    seg {seg_id}: \"{src[:50]}\" → \"{tgt[:50]}\"")

    if not dry_run:
        text_to_id = {r["source_text"]: r["id"] for r in rows}

        # Update segments that came from TM
        for src_text, tgt_text in tm_hits.items():
            seg_id = text_to_id.get(src_text)
            if seg_id:
                cur.execute("""
                    UPDATE segments
                    SET translated_text    = %s,
                        translation_source = 'TM_EXACT'::translation_source
                    WHERE id = %s;
                """, (tgt_text, seg_id))

        # Update segments translated by Azure + upsert into TM
        for seg_id, tgt_text in azure_map.items():
            cur.execute("""
                UPDATE segments
                SET translated_text    = %s,
                    translation_source = 'LLM'::translation_source
                WHERE id = %s;
            """, (tgt_text, seg_id))

            src_text = next((r["source_text"] for r in rows if r["id"] == seg_id), None)
            if src_text:
                h = sha256(src_text)
                _upsert_tm_entry(cur, org_id, source_lang, target_lang,
                                 src_text, tgt_text, h, domain)

        conn.commit()
        print("\n  segments.translated_text populated.")
        print("  tm_entries upserted for new translations.")

    print(f"\n  TM_EXACT  : {len(tm_hits)}")
    print(f"  Azure MT  : {len(azure_map)}")
    print(f"  Success   : {len(tm_hits) + len(azure_map)} / {total}")

    cur.close()
    conn.close()


# ══════════════════════════════════════════════════════════════════════════════
# STEP 4 — BACK-TRANSLATION QUALITY AGENT
# Schema table used: agent_segment_evaluation (already defined in schema — no CREATE needed)
#
# agent_segment_evaluation columns written:
#   segment_id, document_id, project_id, source_text, translated_text,
#   back_translated_text, source_lang, target_lang, domain,
#   semantic_score, glossary_score, tm_score, final_sqs,
#   decision, needs_linguist_review
# ══════════════════════════════════════════════════════════════════════════════

def _mymemory_translate(text: str, src: str, tgt: str) -> str:
    if not text or not text.strip():
        return ""
    params = {"q": text, "langpair": f"{src}|{tgt}"}
    if MYMEMORY_EMAIL:
        params["de"] = MYMEMORY_EMAIL
    try:
        resp = requests.get(MYMEMORY_URL, params=params, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        if data.get("responseStatus") == 200:
            t = data["responseData"]["translatedText"]
            return (t.replace("&amp;", "&").replace("&lt;", "<")
                     .replace("&gt;", ">").replace("&quot;", '"').replace("&#39;", "'"))
        else:
            print(f"  [MyMemory] status={data.get('responseStatus')} — {data.get('responseDetails')}")
            return ""
    except Exception as e:
        print(f"  [MyMemory] Error: {e}")
        return ""


def run_back_translation(
    document_id: int,
    project_id:  int,
    org_id:      int,
    target_lang: str,
):
    """
    Back-translate each translated segment → English via MyMemory,
    score quality (semantic + glossary + TM), store in agent_segment_evaluation.

    Reads from : segments (translated_text), glossary_terms, tm_entries
    Writes to  : agent_segment_evaluation
    """
    step_banner(4, "BACK-TRANSLATION QUALITY AGENT")

    # Heavy imports — loaded only when this step runs
    from sentence_transformers import SentenceTransformer
    from sklearn.metrics.pairwise import cosine_similarity as cos_sim
    from rapidfuzz import fuzz

    print("  Loading sentence embedding model …")
    model = SentenceTransformer("all-MiniLM-L6-v2")

    conn = get_conn()
    cur  = conn.cursor()

    # Load translated segments
    cur.execute("""
        SELECT id, source_text, translated_text
        FROM   segments
        WHERE  document_id     = %s
          AND  translated_text IS NOT NULL
        ORDER  BY id;
    """, (document_id,))
    rows = cur.fetchall()

    # Load glossary terms (EN → target_lang) for glossary scoring
    cur.execute("""
        SELECT source_term, target_term
        FROM   glossary_terms
        WHERE  org_id      = %s
          AND  source_lang = 'en'
          AND  target_lang = %s;
    """, (org_id, target_lang))
    glossary = cur.fetchall()

    # Load TM entries for fuzzy TM scoring
    cur.execute("""
        SELECT source_text, target_text
        FROM   tm_entries
        WHERE  org_id      = %s
          AND  source_lang = 'en'
          AND  target_lang = %s;
    """, (org_id, target_lang))
    tm_entries = cur.fetchall()

    total = len(rows)
    print(f"  Segments to evaluate : {total}")
    print(f"  Glossary terms       : {len(glossary)}")
    print(f"  TM entries           : {len(tm_entries)}\n")

    # ── Scoring helpers ────────────────────────────────────────────────────────

    def compute_semantic(orig: str, back: str) -> float:
        e1 = model.encode([orig])
        e2 = model.encode([back])
        return float(cos_sim(e1, e2)[0][0])

    def compute_glossary(orig: str, translated: str, back: str) -> float:
        total_g, correct = 0, 0
        for eng, tgt in glossary:
            if eng.lower() in orig.lower():
                total_g += 1
                if tgt.lower() in translated.lower() and eng.lower() in back.lower():
                    correct += 1
        return 1.0 if total_g == 0 else correct / total_g

    def compute_tm_score(orig: str) -> float:
        if not tm_entries:
            return 0.0
        best = max(fuzz.ratio(orig, src) for src, _ in tm_entries)
        return best / 100.0

    def compute_sqs(S: float, G: float, T: float) -> float:
        return 0.5 * S + 0.3 * G + 0.2 * T

    def make_decision(sqs: float) -> str:
        if sqs >= 0.85: return "ACCEPT"
        if sqs >= 0.70: return "REVIEW"
        if sqs >= 0.50: return "RETRANSLATE"
        return "LINGUIST REVIEW"

    # ── Process each segment ───────────────────────────────────────────────────

    results_summary = []

    for seg_id, source_text, translated_text in rows:
        print(f"  ─── seg {seg_id} ─────────────────────────────────────────")
        print(f"  Source (EN)    : {source_text}")
        print(f"  Translated     : {translated_text}")

        back_text = _mymemory_translate(translated_text, target_lang, "en-gb")
        print(f"  Back-transl    : {back_text}")

        if not back_text:
            print("  ⚠️  Back-translation failed — skipping.\n")
            continue

        S   = compute_semantic(source_text, back_text)
        G   = compute_glossary(source_text, translated_text, back_text)
        T   = compute_tm_score(source_text)
        sqs = compute_sqs(S, G, T)
        dec = make_decision(sqs)

        print(f"  Semantic score : {round(S, 3)}")
        print(f"  Glossary score : {round(G, 3)}")
        print(f"  TM score       : {round(T, 3)}")
        print(f"  Final SQS      : {round(sqs, 3)}")
        print(f"  Decision       : {dec}\n")

        needs_review = dec in ("REVIEW", "LINGUIST REVIEW")

        # agent_segment_evaluation is already defined in schema — no CREATE TABLE needed
        cur.execute("""
            INSERT INTO agent_segment_evaluation (
                segment_id, document_id, project_id,
                source_text, translated_text, back_translated_text,
                source_lang, target_lang, domain,
                semantic_score, glossary_score, tm_score, final_sqs,
                decision, needs_linguist_review
            ) VALUES (%s, %s, %s, %s, %s, %s,
                      'en', %s, 'general',
                      %s, %s, %s, %s,
                      %s, %s);
        """, (
            seg_id, document_id, project_id,
            source_text, translated_text, back_text,
            target_lang,
            float(S), float(G), float(T), float(sqs),
            dec, needs_review,
        ))
        conn.commit()

        results_summary.append({"seg_id": seg_id, "sqs": sqs, "decision": dec})

    cur.close()
    conn.close()

    # ── Summary ────────────────────────────────────────────────────────────────
    if results_summary:
        avg_sqs   = sum(r["sqs"] for r in results_summary) / len(results_summary)
        decisions = {}
        for r in results_summary:
            decisions[r["decision"]] = decisions.get(r["decision"], 0) + 1

        print("  ═══════════════════════════════════════════════════════════")
        print("  BACK-TRANSLATION SUMMARY")
        print("  ═══════════════════════════════════════════════════════════")
        print(f"  Evaluated       : {len(results_summary)} / {total}")
        print(f"  Average SQS     : {round(avg_sqs, 3)}")
        for dec, count in sorted(decisions.items()):
            print(f"  {dec:<22}: {count}")
        print("  Results stored  : agent_segment_evaluation table")


# ══════════════════════════════════════════════════════════════════════════════
# MASTER PIPELINE
# ══════════════════════════════════════════════════════════════════════════════

def run_pipeline(
    input_path:      Path,
    project_id:      int,
    target_lang:     str,
    org_id:          int,
    source_lang:     str           = "en",
    domain:          Optional[str] = None,
    dry_run:         bool          = False,
    validation_lang: str           = "en-US",
):
    banner("VaaniSetu — Full Translation Pipeline")
    print(f"  Input file   : {input_path}")
    print(f"  Project ID   : {project_id}")
    print(f"  Org ID       : {org_id}")
    print(f"  Source lang  : {source_lang}")
    print(f"  Target lang  : {target_lang}")
    print(f"  Domain       : {domain or 'general'}")
    print(f"  Dry run      : {dry_run}")

    t0 = time.time()

    # Step 1 — Parsing
    document_id = run_parsing(input_path, project_id, target_lang)

    # Step 2 — Validation
    issues = run_validation(document_id, lang=validation_lang)
    if issues:
        print(f"\n  [info] {len(issues)} segment(s) had source-text issues.")
        print("  [info] Pipeline continues — issues are logged but not blocking.")

    # Step 3 — Translation
    run_translation(
        document_id = document_id,
        target_lang = target_lang,
        source_lang = source_lang,
        org_id      = org_id,
        domain      = domain,
        dry_run     = dry_run,
    )

    # Step 4 — Back-translation quality check
    if dry_run:
        print("\n  [dry-run] Skipping back-translation (no DB translations to evaluate).")
    else:
        run_back_translation(
            document_id = document_id,
            project_id  = project_id,
            org_id      = org_id,
            target_lang = target_lang,
        )

    elapsed = round(time.time() - t0, 1)
    banner("PIPELINE COMPLETE")
    print(f"  Document ID  : {document_id}")
    print(f"  Elapsed      : {elapsed}s")
    print(f"  DB tables written:")
    print(f"    • documents                  (1 row)")
    print(f"    • segments                   (parsed + translated chunks)")
    print(f"    • validation_issues          (source errors, if any)")
    print(f"    • tm_entries                 (new translations cached)")
    print(f"    • agent_segment_evaluation   (back-translation scores)")
    print()


# ══════════════════════════════════════════════════════════════════════════════
# CLI  —  only --input is required; everything else has sensible defaults
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="VaaniSetu end-to-end translation pipeline"
    )
    parser.add_argument("--input",           required=True,
                        help="Path to the .docx input file")
    parser.add_argument("--project-id",      type=int, default=1,
                        help="Project ID in the projects table (default: 1)")
    parser.add_argument("--target-lang",     default="de",
                        help="Target language code (default: de = German)")
    parser.add_argument("--org-id",          type=int, default=101,
                        help="Organisation ID (default: 101)")
    parser.add_argument("--source-lang",     default="en",
                        help="Source language code (default: en)")
    parser.add_argument("--domain",          default=None,
                        help="Domain tag for TM entries e.g. legal, medical")
    parser.add_argument("--validation-lang", default="en-US",
                        help="LanguageTool language code (default: en-US)")
    parser.add_argument("--dry-run",         action="store_true",
                        help="Parse + validate only, skip translation DB writes")

    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"[ERROR] File not found: {input_path}")
        sys.exit(1)
    if input_path.suffix.lower() != ".docx":
        print(f"[ERROR] Only .docx files are supported. Got: {input_path.suffix}")
        sys.exit(1)

    run_pipeline(
        input_path      = input_path,
        project_id      = args.project_id,
        target_lang     = args.target_lang,
        org_id          = args.org_id,
        source_lang     = args.source_lang,
        domain          = args.domain,
        dry_run         = args.dry_run,
        validation_lang = args.validation_lang,
    )
