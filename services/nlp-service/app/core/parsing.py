import argparse
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

WORDS_PER_SEGMENT = 10
SCHEMA_VERSION = "1.0"

HEADING_STYLES = {
    "heading 1", "heading 2", "heading 3", "heading 4", "heading 5", "heading 6"
}
CAPTION_STYLES = {"caption"}
LIST_STYLES = {
    "list paragraph", "listbullet", "listbullet2", "listnumber", "listnumber2",
    "list bullet", "list number"
}


def _style_id(para: Paragraph) -> str:
    try:
        return para.style.style_id or ""
    except Exception:
        return ""


def _style_name(para: Paragraph) -> str:
    try:
        return (para.style.name or "").lower().strip()
    except Exception:
        return ""


def _alignment(para: Paragraph) -> str:
    alignment_map = {
        "LEFT": "left",
        "CENTER": "center",
        "RIGHT": "right",
        "JUSTIFY": "justify",
        None: "left",
    }
    try:
        val = para.alignment
        return alignment_map.get(str(val).split(".")[-1] if val else None, "left")
    except Exception:
        return "left"


def _is_list_para(para: Paragraph) -> bool:
    if _style_name(para) in LIST_STYLES:
        return True
    p_pr = para._p.find(qn("w:pPr"))
    return p_pr is not None and p_pr.find(qn("w:numPr")) is not None


def _has_inline_image(para: Paragraph) -> bool:
    xml = para._p.xml
    return any(token in xml for token in ("w:drawing", "w:pict", "a:blip", "wp:inline", "wp:anchor"))


def _run_props(run) -> dict:
    props = {
        "run_index": 0,
        "text": run.text,
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline),
    }
    try:
        props["font_name"] = run.font.name or ""
    except Exception:
        props["font_name"] = ""
    try:
        size = run.font.size
        props["font_size_pt"] = int(size.pt) if size else None
    except Exception:
        props["font_size_pt"] = None
    try:
        color = run.font.color.rgb
        props["color_hex"] = str(color) if color else None
    except Exception:
        props["color_hex"] = None
    return props


def _runs_for_para(para: Paragraph) -> list[dict]:
    runs = []
    for index, run in enumerate(para.runs):
        if not run.text:
            continue
        props = _run_props(run)
        props["run_index"] = index
        runs.append(props)
    return runs


def _count_words(text: str) -> int:
    return len(text.split())


def _split_runs_into_chunks(runs: list[dict], words_per_segment: int) -> list[list[dict]]:
    if not runs:
        return []

    chunks: list[list[dict]] = []
    current_chunk: list[dict] = []
    current_words = 0

    for run in runs:
        words_in_run = _count_words(run["text"])

        if current_words + words_in_run > words_per_segment and current_chunk:
            chunks.append(current_chunk)
            current_chunk = []
            current_words = 0

        if words_in_run > words_per_segment:
            words = run["text"].split()
            for start in range(0, len(words), words_per_segment):
                piece = words[start:start + words_per_segment]
                text = " ".join(piece)
                chunks.append([{**run, "text": text, "run_index": run["run_index"]}])
        else:
            current_chunk.append(run)
            current_words += words_in_run

    if current_chunk:
        chunks.append(current_chunk)

    return chunks


def _chunk_source_text(runs: list[dict]) -> str:
    return "".join(run["text"] for run in runs).strip()


class SegmentBuilder:
    def __init__(self, words_per_segment: int):
        self.words_per_segment = words_per_segment
        self._counter = 0
        self.segments: list[dict] = []
        self.non_translatable: list[dict] = []

    def _next_id(self) -> str:
        self._counter += 1
        return f"seg_{self._counter:03d}"

    def add_para(self, para: Paragraph, para_idx: int, seg_type: str, extra_anchor: dict | None = None):
        runs = _runs_for_para(para)
        full_text = para.text.strip()

        if not full_text:
            self.non_translatable.append({
                "type": "EMPTY_PARAGRAPH",
                "position_after_segment": f"seg_{self._counter:03d}",
                "docx_anchor": {
                    "xpath": f"/w:document/w:body/w:p[{para_idx + 1}]",
                    "paragraph_index": para_idx,
                },
            })
            return

        if _has_inline_image(para):
            nt = {
                "type": "IMAGE",
                "docx_anchor": {
                    "xpath": f"/w:document/w:body/w:p[{para_idx + 1}]",
                    "paragraph_index": para_idx,
                },
            }
            if self._counter:
                nt["position_after_segment"] = f"seg_{self._counter:03d}"
            self.non_translatable.append(nt)

        chunks = _split_runs_into_chunks(runs, self.words_per_segment)
        if not chunks and full_text:
            chunks = [[{
                "run_index": 0,
                "text": full_text,
                "bold": False,
                "italic": False,
                "underline": False,
                "font_name": "",
                "font_size_pt": None,
                "color_hex": None,
            }]]

        for chunk_runs in chunks:
            source_text = _chunk_source_text(chunk_runs)
            if not source_text:
                continue

            seg_id = self._next_id()
            anchor = {
                "xpath": f"/w:document/w:body/w:p[{para_idx + 1}]",
                "paragraph_index": para_idx,
                "style_id": _style_id(para),
                "alignment": _alignment(para),
                "runs": chunk_runs,
            }

            if seg_type == "LIST_ITEM":
                p_pr = para._p.find(qn("w:pPr"))
                if p_pr is not None:
                    num_pr = p_pr.find(qn("w:numPr"))
                    if num_pr is not None:
                        ilvl = num_pr.find(qn("w:ilvl"))
                        num_id = num_pr.find(qn("w:numId"))
                        anchor["list_level"] = int(ilvl.get(qn("w:val"), 0)) if ilvl is not None else 0
                        anchor["numbering_id"] = int(num_id.get(qn("w:val"), 0)) if num_id is not None else 0

            if extra_anchor:
                anchor.update(extra_anchor)

            self.segments.append({
                "id": seg_id,
                "position": self._counter,
                "type": seg_type,
                "source_text": source_text,
                "word_count": _count_words(source_text),
                "is_translatable": True,
                "docx_anchor": anchor,
            })

    def add_table(self, table: Table, table_idx: int):
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                for para in cell.paragraphs:
                    cell_text = para.text.strip()
                    if not cell_text:
                        continue

                    runs = _runs_for_para(para)
                    chunks = _split_runs_into_chunks(runs, self.words_per_segment)
                    if not chunks:
                        chunks = [[{
                            "run_index": 0,
                            "text": cell_text,
                            "bold": False,
                            "italic": False,
                            "underline": False,
                            "font_name": "",
                            "font_size_pt": None,
                            "color_hex": None,
                        }]]

                    for chunk_runs in chunks:
                        source_text = _chunk_source_text(chunk_runs)
                        if not source_text:
                            continue
                        seg_id = self._next_id()
                        self.segments.append({
                            "id": seg_id,
                            "position": self._counter,
                            "type": "TABLE_CELL",
                            "source_text": source_text,
                            "word_count": _count_words(source_text),
                            "is_translatable": True,
                            "docx_anchor": {
                                "xpath": f"/w:document/w:body/w:tbl[{table_idx + 1}]/w:tr[{row_idx + 1}]/w:tc[{col_idx + 1}]/w:p[1]",
                                "table_index": table_idx,
                                "row": row_idx,
                                "col": col_idx,
                                "runs": chunk_runs,
                            },
                        })

    def add_footnotes(self, docx_path: Path):
        try:
            with zipfile.ZipFile(docx_path) as zf:
                if "word/footnotes.xml" not in zf.namelist():
                    return
                xml_bytes = zf.read("word/footnotes.xml")
        except Exception:
            return

        from lxml import etree

        root = etree.fromstring(xml_bytes)
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        for footnote in root.findall(f"{{{namespace}}}footnote"):
            footnote_id = footnote.get(f"{{{namespace}}}id", "")
            footnote_type = footnote.get(f"{{{namespace}}}type", "")
            if footnote_type in ("separator", "continuationSeparator"):
                continue

            for paragraph in footnote.findall(f".//{{{namespace}}}p"):
                runs_out = []
                text_parts = []
                for index, run_el in enumerate(paragraph.findall(f"{{{namespace}}}r")):
                    run_text = "".join(t.text or "" for t in run_el.findall(f"{{{namespace}}}t"))
                    if not run_text:
                        continue
                    text_parts.append(run_text)
                    runs_out.append({
                        "run_index": index,
                        "text": run_text,
                        "bold": False,
                        "italic": False,
                        "underline": False,
                        "font_name": "",
                        "font_size_pt": None,
                        "color_hex": None,
                    })

                full_text = "".join(text_parts).strip()
                if not full_text:
                    continue

                for chunk_runs in _split_runs_into_chunks(runs_out, self.words_per_segment) or [[{"run_index": 0, "text": full_text, "bold": False, "italic": False, "underline": False, "font_name": "", "font_size_pt": None, "color_hex": None}]]:
                    source_text = _chunk_source_text(chunk_runs)
                    if not source_text:
                        continue
                    seg_id = self._next_id()
                    self.segments.append({
                        "id": seg_id,
                        "position": self._counter,
                        "type": "FOOTNOTE",
                        "source_text": source_text,
                        "word_count": _count_words(source_text),
                        "is_translatable": True,
                        "docx_anchor": {
                            "source_file": "word/footnotes.xml",
                            "footnote_id": int(footnote_id) if footnote_id.lstrip("-").isdigit() else footnote_id,
                            "xpath": f"/w:footnotes/w:footnote[@w:id='{footnote_id}']/w:p[1]",
                            "runs": chunk_runs,
                        },
                    })

    def scan_page_fields(self, docx_path: Path):
        try:
            with zipfile.ZipFile(docx_path) as zf:
                header_footer_files = [name for name in zf.namelist() if re.match(r"word/(header|footer)\d*\.xml", name)]
                for item in header_footer_files:
                    content = zf.read(item).decode("utf-8", errors="replace")
                    if "PAGE" in content and ("w:fldChar" in content or "w:instrText" in content):
                        self.non_translatable.append({
                            "type": "PAGE_NUMBER",
                            "docx_anchor": {"source_file": item, "field_type": "PAGE"},
                        })
                        break
        except Exception:
            pass


def _iter_body_elements(doc: Document):
    para_index = 0
    table_index = 0
    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            yield "para", Paragraph(child, doc), para_index
            para_index += 1
        elif tag == "tbl":
            yield "table", Table(child, doc), table_index
            table_index += 1


def parse_docx(input_path: Path, doc_id: str, org_id: str, words_per_segment: int = WORDS_PER_SEGMENT, raw_key: str | None = None, source_lang: str = "en-US") -> dict:
    doc = Document(str(input_path))
    builder = SegmentBuilder(words_per_segment)

    for element_type, obj, index in _iter_body_elements(doc):
        if element_type == "para":
            style_name = _style_name(obj)
            if style_name in HEADING_STYLES or style_name.startswith("heading"):
                builder.add_para(obj, index, "HEADING")
            elif style_name in CAPTION_STYLES:
                builder.add_para(obj, index, "CAPTION")
            elif _is_list_para(obj):
                builder.add_para(obj, index, "LIST_ITEM")
            else:
                builder.add_para(obj, index, "PARAGRAPH")
        else:
            builder.add_table(obj, index)

    builder.add_footnotes(input_path)
    builder.scan_page_fields(input_path)

    page_count = 1
    try:
        page_breaks = sum(
            1
            for paragraph in doc.paragraphs
            for run in paragraph.runs
            if run._r.xml and "w:lastRenderedPageBreak" in run._r.xml
        )
        if page_breaks:
            page_count = page_breaks + 1
    except Exception:
        pass

    total_words = sum(segment["word_count"] for segment in builder.segments)

    return {
        "document_id": doc_id,
        "org_id": org_id,
        "source_format": "docx",
        "source_lang": source_lang,
        "schema_version": SCHEMA_VERSION,
        "minio_raw_key": raw_key or input_path.name,
        "word_count": total_words,
        "segment_count": len(builder.segments),
        "page_count": page_count,
        "parsed_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "warnings": [],
        "segments": builder.segments,
        "non_translatable": builder.non_translatable,
    }


def build_segment_rows(structure: dict, document_id: int, target_lang: str) -> list[dict]:
    rows = []
    for segment in structure.get("segments", []):
        source_text = (segment.get("source_text") or "").strip()
        if not source_text:
            continue
        rows.append({
            "document_id": document_id,
            "target_lang": target_lang,
            "source_text": source_text,
            "translated_text": None,
            "translation_source": None,
            "locked_by": None,
        })
    return rows


def main():
    parser = argparse.ArgumentParser(description="Parse a DOCX file into structure.json")
    parser.add_argument("--input", required=True, help="Path to input DOCX")
    parser.add_argument("--output", default="structure.json", help="Output JSON path")
    parser.add_argument("--doc-id", default=None, help="Document ID")
    parser.add_argument("--org-id", default="default-org", help="Organisation ID")
    parser.add_argument("--source-lang", default="en-US", help="Source language")
    parser.add_argument("--raw-key", default=None, help="Original MinIO key")
    parser.add_argument("--words-per-segment", type=int, default=WORDS_PER_SEGMENT)
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    document_id = args.doc_id or f"doc_{uuid.uuid4().hex[:6]}"
    structure = parse_docx(
        input_path=input_path,
        doc_id=document_id,
        org_id=args.org_id,
        words_per_segment=args.words_per_segment,
        raw_key=args.raw_key,
        source_lang=args.source_lang,
    )

    output_path = Path(args.output)
    output_path.write_text(json.dumps(structure, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[parser] Output written to {output_path}")
    print(f"[parser] Segments: {structure['segment_count']}")


if __name__ == "__main__":
    main()
